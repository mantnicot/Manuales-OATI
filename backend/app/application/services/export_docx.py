from __future__ import annotations

import io
from typing import Any

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from app.application.services.document_html import (
    OATI_STATIC,
    _crest_scale_pct,
    _format_figure_caption,
    _image_scale_pct,
    _resolve_image_data,
    oati_image_index_by_block_id,
    resolve_crest_src,
    resolve_logo_oati_src,
    resolve_logo_ud_src,
)
from app.application.services.export_html_images import (
    cleanup_temp_paths,
    materialize_data_images_in_html,
)
from app.application.services.export_media import image_bytes_from_src
from app.application.services.plain_text import plain_from_html
from app.domain.entities.manual import Block, Manual


def _cell_add_logo(cell, src: str | None, placeholder: str, max_w_in: float = 1.15) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    raw = (src or "").strip()
    data = image_bytes_from_src(raw) if raw else None
    if data:
        try:
            run = p.add_run()
            run.add_picture(io.BytesIO(data), width=Inches(max_w_in))
            return
        except Exception:
            pass
    r = p.add_run(placeholder)
    r.font.size = Pt(8)


def _add_oati_header_table(container: Any, d: dict, *, add_spacer_paragraph: bool = True) -> None:
    """Cabecera tabular (4 columnas). Puede añadirse al documento o al encabezado de sección (se repite en cada página)."""
    try:
        table = container.add_table(rows=1, cols=4)
    except TypeError:
        # Encabezado de sección (python-docx): ancho total obligatorio
        table = container.add_table(rows=1, cols=4, width=Inches(7.4))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    try:
        table.style = "Table Grid"
    except Exception:
        pass

    cells = table.rows[0].cells
    w = [Inches(1.35), Inches(2.65), Inches(2.05), Inches(1.35)]
    for c, width in zip(cells, w, strict=False):
        c.width = width

    _cell_add_logo(
        cells[0],
        resolve_logo_ud_src(d),
        "Logo\nUniversidad",
    )

    mid = cells[1]
    mid.text = ""
    p_mid = mid.paragraphs[0]
    r0 = p_mid.add_run(str(d.get("moduleHeaderLine") or ""))
    r0.bold = True
    r0.font.size = Pt(10)
    for extra in (str(d.get("macroProcess") or ""), str(d.get("processLine") or "")):
        if extra.strip():
            p_e = mid.add_paragraph()
            p_e.add_run(extra).font.size = Pt(10)

    meta = cells[2]
    meta.text = ""
    code = str(d.get("code") or "").strip() or "—"
    ver = str(d.get("version") or "").strip() or "—"
    fap = str(d.get("approvalDate") or "").strip() or "—"
    for i, (label, val) in enumerate(
        (("Código:", code), ("Versión:", ver), ("Fecha de Aprobación:", fap)),
    ):
        p = meta.paragraphs[0] if i == 0 else meta.add_paragraph()
        lb = p.add_run(f"{label} ")
        lb.bold = True
        lb.font.size = Pt(9)
        p.add_run(val).font.size = Pt(9)

    _cell_add_logo(
        cells[3],
        resolve_logo_oati_src(d),
        "Logo\nOATI",
    )

    if add_spacer_paragraph:
        container.add_paragraph()


def _append_word_field(paragraph, code: str) -> None:
    """Inserta un campo de Word (PAGE, NUMPAGES, etc.)."""
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = code
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def _footer_with_page_fields(section) -> None:
    """Pie con 'Página X de Y' actualizado por Word y leyenda OATI."""
    foot = section.footer
    foot.is_linked_to_previous = False
    p = foot.paragraphs[0] if foot.paragraphs else foot.add_paragraph()
    p.text = ""
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.55), WD_TAB_ALIGNMENT.RIGHT)
    r0 = p.add_run("Página ")
    r0.font.size = Pt(9)
    r0.font.color.rgb = RGBColor(0x75, 0x85, 0x99)
    _append_word_field(p, "PAGE")
    r_mid = p.add_run(" de ")
    r_mid.font.size = Pt(9)
    r_mid.font.color.rgb = RGBColor(0x75, 0x85, 0x99)
    _append_word_field(p, "NUMPAGES")
    p.add_run("\t")
    r1 = p.add_run("Desarrollado por Oficina Asesora de Sistemas OATI")
    r1.font.size = Pt(7)
    r1.font.color.rgb = RGBColor(0x9c, 0xa3, 0xaf)


def _gap_paragraph(doc: Document, after_pt: float = 14, before_pt: float = 0) -> None:
    g = doc.add_paragraph()
    g.paragraph_format.space_after = Pt(after_pt)
    g.paragraph_format.space_before = Pt(before_pt)


def _add_toc_page(doc: Document, d: dict) -> None:
    """Índice fijo como en la vista HTML (hoja 2); el pie va en el footer de sección."""
    _add_oati_header_table(doc, d)
    _gap_paragraph(doc, after_pt=8, before_pt=6)
    h = doc.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h.paragraph_format.space_before = Pt(10)
    h.paragraph_format.space_after = Pt(18)
    r = h.add_run("TABLA DE CONTENIDO")
    r.bold = True
    r.font.size = Pt(12)
    entries = (
        "INTRODUCCIÓN",
        "1. OBJETIVO",
        "2. ALCANCE",
        "3. RESPONSABLES",
        "4. DEFINICIONES Y SIGLAS",
        "5. DESCRIPCIÓN DE CADA PASO (DETALLADO)",
    )
    for title in entries:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(10)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(6.0), WD_TAB_ALIGNMENT.RIGHT)
        t_run = p.add_run(title)
        t_run.bold = True
        t_run.font.size = Pt(11)
        p.add_run("\t3")


def _add_cover_sheet(doc: Document, d: dict) -> None:
    _add_oati_header_table(doc, d)
    _gap_paragraph(doc, after_pt=20, before_pt=4)

    crest_src = resolve_crest_src(d).strip()
    cdata = image_bytes_from_src(crest_src) if crest_src else None
    crest_w = Inches(1.75 * max(50, min(150, _crest_scale_pct(d))) / 100)
    if cdata:
        try:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(20)
            run = p.add_run()
            run.add_picture(io.BytesIO(cdata), width=crest_w)
        except Exception:
            p = doc.add_paragraph("Escudo institucional")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(20)
    else:
        p = doc.add_paragraph("Escudo institucional")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(20)

    for i, (line, bold, pts) in enumerate(
        (
            (str(d.get("institutionLine") or ""), True, 15),
            (str(d.get("centralTitle") or ""), True, 13),
            (str(d.get("unitFooterLine") or ""), True, 11),
        ),
    ):
        p = doc.add_paragraph(line)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(10 if i > 0 else 6)
        p.paragraph_format.space_after = Pt(12 if i < 2 else 6)
        if p.runs:
            p.runs[0].bold = bold
            p.runs[0].font.size = Pt(pts)


def _add_centered_figure_word(doc: Document, src: str, caption: str, width_pct: int) -> None:
    data = image_bytes_from_src(src)
    pct = max(25, min(100, int(width_pct)))
    w_in = 6.5 * pct / 100
    if data:
        try:
            pic_stream = io.BytesIO(data)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run()
            run.add_picture(pic_stream, width=Inches(w_in))
        except Exception:
            doc.add_paragraph(
                "[No se pudo insertar la imagen en Word; pruebe PNG/JPEG o vuelva a guardar el manual.]",
                style=None,
            )
    elif (src or "").strip():
        doc.add_paragraph(
            "(Sin datos válidos de imagen para Word; compruebe la URL o suba el archivo de nuevo y guarde.)",
            style=None,
        )
    if caption:
        cp = doc.add_paragraph(caption)
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in cp.runs:
            r.bold = True


def _add_oati_section_block(
    doc: Document,
    b: Block | None,
    heading: str,
    img_idx: dict[str, int],
) -> None:
    if not b:
        return
    d = b.data
    doc.add_heading(heading, level=1)
    doc.add_paragraph(plain_from_html(str(d.get("text", ""))))
    src = _resolve_image_data(d)
    if src:
        cap = _format_figure_caption(img_idx.get(b.id), str(d.get("imageCaption", "")))
        _add_centered_figure_word(doc, src, cap, _image_scale_pct(d))


def _add_oati_note_paragraph(doc: Document, body: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    t1 = p.add_run("Nota: ")
    t1.bold = True
    t1.italic = True
    t2 = p.add_run(plain_from_html(body))
    t2.italic = True
    t2.font.underline = True


def _build_oati_document(doc: Document, manual: Manual, img_idx: dict[str, int]) -> None:
    blocks = sorted(manual.blocks, key=lambda x: x.order)
    smap: dict[str, Block] = {b.type: b for b in blocks if b.type in OATI_STATIC}
    cover = smap.get("oati_cover")
    d_cover: dict = cover.data if cover else {}

    sec_cover = doc.sections[0]
    sec_cover.footer.is_linked_to_previous = False
    _add_cover_sheet(doc, d_cover)
    _footer_with_page_fields(sec_cover)

    doc.add_section(WD_SECTION_START.NEW_PAGE)
    sec_toc = doc.sections[-1]
    sec_toc.footer.is_linked_to_previous = False
    _add_toc_page(doc, d_cover)
    _footer_with_page_fields(sec_toc)

    doc.add_section(WD_SECTION_START.NEW_PAGE)
    sec_body = doc.sections[-1]
    sec_body.header.is_linked_to_previous = False
    _add_oati_header_table(sec_body.header, d_cover, add_spacer_paragraph=False)
    _gap_paragraph(doc, after_pt=12, before_pt=2)

    for typ, title in (
        ("oati_intro", "INTRODUCCIÓN"),
        ("oati_objective", "1. OBJETIVO"),
        ("oati_scope", "2. ALCANCE"),
        ("oati_responsible", "3. RESPONSABLES"),
        ("oati_definitions", "4. DEFINICIONES Y SIGLAS"),
    ):
        _add_oati_section_block(doc, smap.get(typ), title, img_idx)

    doc.add_heading("5. DESCRIPCIÓN DE CADA PASO (DETALLADO)", level=1)

    flow = [b for b in blocks if b.type in ("oati_step", "oati_note")]
    step_num = 0
    for b in flow:
        if b.type == "oati_step":
            step_num += 1
            d = b.data
            title = str(d.get("title", "Paso"))
            doc.add_heading(f"5.{step_num} {title.upper()}", level=2)
            doc.add_paragraph(plain_from_html(str(d.get("description", ""))))
            src = _resolve_image_data(d)
            if src:
                cap = _format_figure_caption(img_idx.get(b.id), str(d.get("imageCaption", "")))
                _add_centered_figure_word(doc, src, cap, _image_scale_pct(d))
        elif b.type == "oati_note":
            _add_oati_note_paragraph(doc, str(b.data.get("body", "")))

    sec_body.footer.is_linked_to_previous = False
    _footer_with_page_fields(sec_body)


def _manual_to_docx_native(manual: Manual) -> bytes:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(11)

    ordered = sorted(manual.blocks, key=lambda x: x.order)
    if not ordered:
        doc.add_paragraph(
            "(Sin bloques en el manual. Guarde el contenido en el editor y vuelva a exportar.)",
        )
    elif manual.blocks and any(b.type == "oati_cover" for b in manual.blocks):
        img_idx = oati_image_index_by_block_id(manual)
        _build_oati_document(doc, manual, img_idx)
    else:
        for b in ordered:
            _append_block(doc, b)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def manual_to_docx(manual: Manual) -> bytes:
    """Plantilla OATI: python-docx estructurado (coherente con la vista previa). Otros: HTML→htmldocx."""
    is_oati = bool(manual.blocks) and any(b.type == "oati_cover" for b in manual.blocks)
    if is_oati:
        return _manual_to_docx_native(manual)
    from app.application.services.document_html import blocks_to_html  # noqa: PLC0415

    html = blocks_to_html(manual)
    html_local, temp_files = materialize_data_images_in_html(html)
    try:
        from htmldocx import HtmlToDocx  # noqa: PLC0415

        document = HtmlToDocx().parse_html_string(html_local)
        document.styles["Normal"].font.name = "Arial"
        document.styles["Normal"].font.size = Pt(11)
        buf = io.BytesIO()
        document.save(buf)
        return buf.getvalue()
    except Exception:
        return _manual_to_docx_native(manual)
    finally:
        cleanup_temp_paths(temp_files)


def _append_block(doc: Document, b: Block) -> None:
    d = b.data
    t = b.type
    if t == "cover":
        for line in [
            d.get("institution", ""),
            d.get("unit", ""),
            d.get("manualTitle", ""),
            d.get("moduleName", ""),
            f"Versión {d.get('version','')} — {d.get('date','')}",
        ]:
            p = doc.add_paragraph(str(line))
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_page_break()
        return
    if t == "toc":
        doc.add_heading(str(d.get("title", "Tabla de contenido")), level=1)
        for e in d.get("entries") or []:
            doc.add_paragraph(str(e.get("text", "")), style="List Bullet")
        return
    if t == "heading1":
        doc.add_heading(str(d.get("text", "")), level=1)
        return
    if t == "heading2":
        doc.add_heading(str(d.get("text", "")), level=2)
        return
    if t == "heading3":
        doc.add_heading(str(d.get("text", "")), level=3)
        return
    if t == "rich_text":
        from html.parser import HTMLParser  # noqa: PLC0415

        class TP(HTMLParser):
            def __init__(self):
                super().__init__()
                self.parts: list[str] = []

            def handle_data(self, data):
                self.parts.append(data)

        p = TP()
        p.feed(str(d.get("html", "")))
        doc.add_paragraph("".join(p.parts))
        return
    if t == "steps":
        for it in d.get("items") or []:
            doc.add_paragraph(str(it.get("title", "")), style="List Number")
            doc.add_paragraph(str(it.get("body", "")))
        return
    if t == "table":
        rows = d.get("rows") or []
        if not rows:
            return
        table = doc.add_table(rows=len(rows), cols=len(rows[0]))
        for r_i, row in enumerate(rows):
            for c_i, cell in enumerate(row):
                table.rows[r_i].cells[c_i].text = str(cell)
        return
    if t == "note":
        p = doc.add_paragraph(str(d.get("text", d.get("html", ""))))
        p.paragraph_format.left_indent = Pt(12)
        return
    if t == "separator":
        doc.add_paragraph("—" * 20)
        return
    if t == "link":
        doc.add_paragraph(f"{d.get('label','')}: {d.get('href','')}")
        return
    if t == "footer":
        doc.add_paragraph(
            f"{d.get('left','')} | {d.get('center','')} | {d.get('right','')}",
        )
